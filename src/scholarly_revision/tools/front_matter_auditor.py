'''Detect configurable front-matter placeholders and template remnants.'''
from __future__ import annotations
import re
from pathlib import Path
from docx import Document
from scholarly_revision.models.scientific_audit import FrontMatterAuditResult
from scholarly_revision.tools._audit_common import IssueFactory,structure_from

DEFAULT_PLACEHOLDERS=[
 r'\bFirstName\b',r'\bLastName\b',r'\bAuthor Name\b',r'\bAffiliation placeholder\b',
 r'\bemail@example\.com\b',r'\bold manuscript title\b',r'\bunrelated DOI\b',
 r'\bold (?:volume|issue|publication date)\b',r'\b(?:submission|acceptance) (?:date|placeholder)\b',
]

def audit_front_matter(source,*,config:dict|None=None)->FrontMatterAuditResult:
    structure=structure_from(source); make=IssueFactory('FRONT_MATTER','FRONT'); issues=[]
    patterns=list((config or {}).get('placeholder_patterns',DEFAULT_PLACEHOLDERS))
    boundary=next((e.order_index for e in structure.elements if e.element_type=='heading' and e.text.strip().casefold() in
        {'abstract','introduction','background'}),len(structure.elements))
    front=[e for e in structure.elements if e.order_index<boundary]
    for element in front:
        for pattern in patterns:
            match=re.search(pattern,element.text,re.I)
            if match:
                issues.append(make('MAJOR','Potential front-matter template placeholder detected.',
                    element=element,evidence=[f'exact match: {match.group(0)}',f'pattern: {pattern}'],manual=True))
    title_texts=[e.text.strip().casefold() for e in front if e.text.strip() and
        (e.text.strip()==(structure.title or '').strip() or e.element_type=='heading')]
    if len(title_texts)!=len(set(title_texts)):
        issues.append(make('MAJOR','Possible duplicated title detected in front matter.',manual=True))
    abstract_count=sum(e.text.strip().casefold()=='abstract' for e in structure.elements)
    if abstract_count>1:
        issues.append(make('MAJOR','Duplicated Abstract heading detected.',evidence=[f'count: {abstract_count}']))
    headers=0;footers=0
    if not hasattr(source,'elements'):
        document=Document(Path(source))
        for section_index,section in enumerate(document.sections):
            for kind,container in [('header',section.header),('footer',section.footer)]:
                if kind=='header':headers+=1
                else:footers+=1
                text=' '.join(p.text for p in container.paragraphs)
                for pattern in [*patterns,*list((config or {}).get('running_header_patterns',[]))]:
                    match=re.search(pattern,text,re.I)
                    if match:
                        issues.append(make('MAJOR',f'Potential unrelated running {kind} or template remnant detected.',
                            evidence=[f'section index: {section_index}',f'exact match: {match.group(0)}'],manual=True))
    return FrontMatterAuditResult(issues=issues,inspected_headers=headers,inspected_footers=footers,
        placeholder_patterns_checked=patterns,checked_element_count=len(front),
        manual_review_required=any(i.manual_review_required for i in issues))

audit_frontmatter=audit_front_matter
